"""처음부터 끝까지 한 사람이 쓰는 흐름 그대로의 통합 테스트.

단위 테스트는 부품이 맞는지 본다. 이 파일은 부품을 **이어 붙였을 때** 맞는지 본다.
실제 사용 순서대로 입력 → 변환 → 수정 → 인사 → 내보내기 → 다시 불러오기 →
비우기 → 되돌리기 → 시점 복구 → 종료 → 재시작을 한 번에 지나간다.
대화상자는 사용자가 누를 버튼을 미리 정해 두고(monkeypatch) 진짜 코드 경로를 탄다.
"""

from __future__ import annotations

import pytest
from PySide6.QtCore import Qt
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QApplication, QDialog, QFileDialog, QMessageBox

from chugui.models import Attendance, Payment, Relation
from chugui.parsing.text_parser import parse_text
from chugui.services.settlement import settle
from chugui.services.text_export import export_text, format_guest_line
from chugui.ui import dialogs as dialogs_module
from chugui.ui.guest_model import Column

ROSTER = (
    "1 홍길동 20만원 친척\n"
    "2 김철수 10만 대학동기\n"
    "3 이영희 30만원 회사 식권2\n"
    "4 박민수 5만원 불참 계좌\n"
    "5 김가족,김친지 30만 이모 소인1\n"
    "6 무명 오만원\n"  # 금액을 읽을 수 없는 줄 → 확인 필요
)


@pytest.fixture
def answers(monkeypatch):
    """대화상자에서 사용자가 누를 답. 테스트 중에 바꿀 수 있다."""
    state = {"question": QMessageBox.StandardButton.Yes, "save_path": ""}
    monkeypatch.setattr(QMessageBox, "question", lambda *a, **k: state["question"])
    monkeypatch.setattr(QMessageBox, "warning", lambda *a, **k: None)
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *a, **k: (state["save_path"], ""))
    return state


@pytest.fixture
def app_window(qt_app, answers):
    from chugui.ui.main_window import MainWindow

    win = MainWindow()
    win.resize(1360, 880)
    win.show()
    qt_app.processEvents()
    yield win
    if win.isVisible():
        win.close()


def _row_of(window, name):
    return next(i for i, g in enumerate(window._model.guests) if g.name == name)


def test_full_user_session(app_window, qt_app, answers, tmp_path):
    w = app_window
    model = w._model

    # 1. 빈 화면: 표 대신 안내.
    assert w._table_stack.currentWidget() is w._empty_state

    # 2. 붙여넣기 → 미리보기가 뜬다.
    w._input.setPlainText(ROSTER)
    w._update_preview()
    assert "6건" in w._preview.text() and "확인 필요 1건" in w._preview.text()

    # 3. Ctrl+Enter 로 변환.
    QTest.keyClick(w, Qt.Key.Key_Return, Qt.KeyboardModifier.ControlModifier)
    qt_app.processEvents()
    assert len(model.guests) == 6
    assert w._table_stack.currentWidget() is w._table
    assert w._review_banner.isVisible()
    couple = model.guests[_row_of(w, "김가족 & 김친지")]
    assert (couple.adult_tickets, couple.child_tickets, couple.relation) == (2, 1, Relation.FAMILY)
    absent = model.guests[_row_of(w, "박민수")]
    assert (absent.attendance, absent.payment, absent.adult_tickets) == (
        Attendance.ABSENT, Payment.TRANSFER, 0,
    )

    # 4. 확인 필요 행 고치기(표에서 금액 수정) → 배너가 사라진다.
    row = _row_of(w, "무명")
    assert model.setData(model.index(row, Column.AMOUNT), "50000", Qt.ItemDataRole.EditRole)
    qt_app.processEvents()
    assert model.guests[row].amount == 50_000
    model.guests[row].warnings.clear()
    model.dataChanged.emit(model.index(row, 0), model.index(row, Column.RAW))
    model.guestsChanged.emit()
    qt_app.processEvents()
    assert not w._review_banner.isVisible()

    # 5. 셀 편집 되돌리기 · 다시 실행.
    hong = _row_of(w, "홍길동")
    model.setData(model.index(hong, Column.AMOUNT), "250000", Qt.ItemDataRole.EditRole)
    assert model.guests[hong].amount == 250_000
    QTest.keyClick(w, Qt.Key.Key_Z, Qt.KeyboardModifier.ControlModifier)
    assert model.guests[hong].amount == 200_000
    QTest.keyClick(w, Qt.Key.Key_Y, Qt.KeyboardModifier.ControlModifier)
    assert model.guests[hong].amount == 250_000

    # 6. 식대 단가 → 순정산이 다시 계산된다.
    w._spin_adult.setValue(50_000)
    w._spin_child.setValue(30_000)
    qt_app.processEvents()
    expected = settle(model.guests, 50_000, 30_000)
    assert w._card_net.value_label.text() == f"{expected.net_amount:,}원"

    # 7. 인사 복사 → 클립보드, 발송 체크는 따로.
    proxy_index = w._proxy.mapFromSource(model.index(hong, Column.COPY))
    w._on_copy_clicked(proxy_index)
    assert "홍길동" in QApplication.clipboard().text()
    assert not model.guests[hong].sent_thanks
    model.setData(model.index(hong, Column.SENT), Qt.CheckState.Checked.value, Qt.ItemDataRole.CheckStateRole)
    assert model.guests[hong].sent_thanks

    # 7-1. 줄 삭제(표에서 Delete) → Ctrl+Z 로 같은 자리에 복원.
    w._table.setFocus()
    lee = _row_of(w, "이영희")
    w._table.selectRow(w._proxy.mapFromSource(model.index(lee, 0)).row())
    QTest.keyClick(w._table, Qt.Key.Key_Delete)
    assert "이영희" not in [g.name for g in model.guests]
    assert [g.guest_id for g in model.guests] == list(range(1, len(model.guests) + 1))
    QTest.keyClick(w, Qt.Key.Key_Z, Qt.KeyboardModifier.ControlModifier)
    assert model.guests[lee].name == "이영희"

    # 8. 검색 · 관계 필터 · 미발송 필터.
    w._search.setText("이영희")
    assert w._proxy.rowCount() == 1
    w._search.clear()
    w._relation_filter.setCurrentText(Relation.FAMILY.value)
    assert w._proxy.rowCount() == 2
    w._relation_filter.setCurrentIndex(0)
    w._chk_unsent.setChecked(True)
    assert w._proxy.rowCount() == 5
    w._chk_unsent.setChecked(False)

    # 9. 추가 모드: 완전 중복(김철수, 이름·금액 동일)은 건너뛰고 새 줄만 붙는다.
    #    홍길동은 5단계에서 금액을 고쳤으므로 여기 쓰면 '중복 의심' 으로 추가되는 게 맞다.
    w._input.setPlainText("2 김철수 10만 대학동기\n7 최신규 7만원 교회")
    QTest.keyClick(w, Qt.Key.Key_Return, Qt.KeyboardModifier.ControlModifier | Qt.KeyboardModifier.ShiftModifier)
    names = [g.name for g in model.guests]
    assert names.count("최신규") == 1 and len(names) == 7

    # 10. 은행 CSV 병합.
    bank = tmp_path / "입출금내역.csv"
    bank.write_text("거래일시,보낸분,입금액\n2026-09-01,정계좌,100000\n", encoding="utf-8-sig")
    w._load_file(str(bank))
    assert model.guests[-1].name == "정계좌"
    assert model.guests[-1].payment is Payment.TRANSFER
    assert len(model.guests) == 8

    # 11. 엑셀로 내보내기 (Ctrl+S) → 엑셀에서 고쳐 다시 올리기 → 교체.
    xlsx = tmp_path / "정산.xlsx"
    answers["save_path"] = str(xlsx)
    QTest.keyClick(w, Qt.Key.Key_S, Qt.KeyboardModifier.ControlModifier)
    assert xlsx.exists()

    from openpyxl import load_workbook

    workbook = load_workbook(xlsx)
    sheet = workbook["축의금 명단"]
    headers = [cell.value for cell in sheet[1]]
    sheet.cell(row=2, column=headers.index("축의금액") + 1).value = 300_000
    workbook.save(xlsx)
    before_total = settle(model.guests).total_amount
    w._load_file(str(xlsx))
    assert len(model.guests) == 8  # 병합이 아니라 교체
    assert model.guests[0].amount == 300_000
    assert model.guests[0].sent_thanks  # 발송 체크도 살아서 돌아온다
    assert settle(model.guests).total_amount == before_total + 50_000

    # 12. 텍스트로 내보내기 → 입력창에 다시 붙여넣으면 같은 명단.
    text_path = tmp_path / "정산.txt"
    answers["save_path"] = str(text_path)
    captured = {}

    class _SaveImmediately(dialogs_module.TextExportDialog):
        def exec(self):
            captured["text"] = self.text
            self._request_save()
            return int(QDialog.DialogCode.Accepted)

    import chugui.ui.main_window as main_module

    original = main_module.TextExportDialog
    main_module.TextExportDialog = _SaveImmediately
    try:
        QTest.keyClick(w, Qt.Key.Key_S, Qt.KeyboardModifier.ControlModifier | Qt.KeyboardModifier.ShiftModifier)
    finally:
        main_module.TextExportDialog = original
    saved = text_path.read_text(encoding="utf-8-sig")
    assert saved == captured["text"]
    assert f"최종 순 정산금 {w._current_settlement().net_amount:,}원" in saved
    reparsed = parse_text(saved)
    assert [(g.names, g.amount, g.relation, g.attendance) for g in reparsed] == [
        (g.names, g.amount, g.relation, g.attendance) for g in model.guests
    ]

    # 13. 전체 비우기 → 되돌리기.
    count = len(model.guests)
    w._handle_clear()
    assert model.guests == [] and w._input.toPlainText() == ""
    w._undo_last()
    assert len(model.guests) == count

    # 14. 이전 시점 복구: 스냅샷 목록에 '전체 비우기 전' 이 있고, 그 시점으로 돌아간다.
    infos = w._snapshots.list_snapshots()
    assert any("전체 비우기 전" in info.label for info in infos)
    model.set_guests(parse_text("임시 1만원"))
    target = next(info for info in infos if "전체 비우기 전" in info.label)
    w._restore_payload(w._snapshots.load(target.path))
    assert len(model.guests) == count

    # 15. 전체 화면 F11 → Esc.
    QTest.keyClick(w, Qt.Key.Key_F11)
    qt_app.processEvents()
    assert w.isFullScreen()
    QTest.keyClick(w, Qt.Key.Key_Escape)
    qt_app.processEvents()
    assert not w.isFullScreen()

    # 16. 종료 → 재시작: 명단 · 입력 · 식대가 그대로.
    snapshot = [(g.name, g.amount, g.sent_thanks) for g in model.guests]
    typed = w._input.toPlainText()
    w.close()
    qt_app.processEvents()

    from chugui.ui.main_window import MainWindow

    again = MainWindow()
    again.show()
    qt_app.processEvents()
    again._restore_session()
    try:
        assert [(g.name, g.amount, g.sent_thanks) for g in again._model.guests] == snapshot
        assert again._input.toPlainText() == typed
        assert (again._spin_adult.value(), again._spin_child.value()) == (50_000, 30_000)
        assert not again._crash_report.crashed  # 정상 종료였다
    finally:
        again.close()


class TestTextExport:
    def test_summary_lines_are_comments(self):
        guests = parse_text("홍길동 10만원 친척")
        text = export_text(guests, settle(guests))
        header = [line for line in text.splitlines() if not line[:1].isdigit()]
        assert all(line.startswith("#") for line in header if line)

    @pytest.mark.parametrize(
        "line",
        ["홍길동 10만원 친척", "박민수 5만원 불참 계좌", "김가족,김친지 30만 이모 소인1",
         "이영희 3만원 교회 식권3", "최동료 10만 A보건지소"],
    )
    def test_line_round_trips(self, line):
        original = parse_text(line)[0]
        back = parse_text(format_guest_line(original))[0]
        assert (back.names, back.amount, back.relation, back.attendance, back.payment,
                back.adult_tickets, back.child_tickets, back.belong) == (
            original.names, original.amount, original.relation, original.attendance,
            original.payment, original.adult_tickets, original.child_tickets, original.belong,
        )

    def test_status_columns_do_not_leak_into_belong(self):
        guest = parse_text("홍길동\t50000\t불참\t계좌\t식권1 소인1")[0]
        assert guest.belong == ""
        assert guest.attendance is Attendance.ABSENT
        assert guest.payment is Payment.TRANSFER

    def test_relation_label_column_is_trusted(self):
        guest = parse_text("박민수  150,000원  종교/모임")[0]
        assert guest.relation is Relation.FAITH
        assert guest.belong == ""
