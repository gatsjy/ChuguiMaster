"""엑셀로 내보내기 → 엑셀에서 수정 → 다시 불러오기.

예전에는 내보낸 파일을 다시 올리면 발송완료 · 참석여부 · 수령경로가 사라지고,
관계는 소속 칸으로 들어갔으며, 목록이 두 배가 되어 전부 '중복 의심' 이 됐다.
"""

from __future__ import annotations

import csv

import pytest

from chugui.models import WARN_NO_AMOUNT, Attendance, Payment, Relation
from chugui.parsing.excel_parser import (
    EXPORT_HEADERS,
    EXPORT_SHEET,
    parse_export_rows,
    parse_rows,
    read_spreadsheet,
)
from chugui.parsing.text_parser import parse_text
from chugui.services.exporter import export_to_excel
from chugui.services.settlement import settle

SOURCE = "홍길동 10만원 친척 식권2\n김철수 5만 불참 계좌\n이영희 3만원 교회 소인1\n김가족,김친지 30만 이모"


@pytest.fixture
def guests():
    items = parse_text(SOURCE)
    items[0].sent_thanks = True
    items[1].note = "나중에 전화"
    return items


@pytest.fixture
def exported(tmp_path, guests):
    path = tmp_path / "명단.xlsx"
    export_to_excel(path, guests, settle(guests))
    return path


def _edit(path, edits):
    """엑셀에서 사용자가 셀을 고친 것처럼 파일을 수정한다."""
    from openpyxl import load_workbook

    workbook = load_workbook(path)
    sheet = workbook[EXPORT_SHEET]
    headers = [cell.value for cell in sheet[1]]
    for (row, title), value in edits.items():
        # cell(..., value=None) 은 값을 지우지 않는다. 직접 대입해야 빈 칸이 된다.
        sheet.cell(row=row + 1, column=headers.index(title) + 1).value = value
    workbook.save(path)


class TestRoundTrip:
    def test_every_field_survives(self, exported, guests):
        result = read_spreadsheet(exported)
        assert result.is_export
        assert len(result.guests) == len(guests)
        for before, after in zip(guests, result.guests, strict=True):
            assert after.name == before.name
            assert after.names == before.names
            assert after.amount == before.amount
            assert after.relation is before.relation
            assert after.attendance is before.attendance
            assert after.payment is before.payment
            assert (after.adult_tickets, after.child_tickets) == (before.adult_tickets, before.child_tickets)
            assert after.belong == before.belong
            assert after.note == before.note
            assert after.sent_thanks is before.sent_thanks

    def test_settlement_is_identical(self, exported, guests):
        assert settle(read_spreadsheet(exported).guests) == settle(guests)

    def test_edits_made_in_excel_are_applied(self, exported):
        _edit(exported, {
            (1, "축의금액"): 150_000,
            (2, "참석여부"): "참석",
            (2, "대인식권"): 2.0,  # 엑셀에서 고치면 실수로 올 수 있다
            (3, "관계분류"): "학교/동창",
            (3, "발송완료"): "완료",
        })
        back = read_spreadsheet(exported).guests
        assert back[0].amount == 150_000
        assert back[1].attendance is Attendance.PRESENT
        assert back[1].adult_tickets == 2
        assert back[2].relation is Relation.SCHOOL
        assert back[2].sent_thanks is True

    def test_rows_deleted_or_added_in_excel(self, exported):
        from openpyxl import load_workbook

        workbook = load_workbook(exported)
        sheet = workbook[EXPORT_SHEET]
        sheet.delete_rows(3)  # 김철수 삭제
        sheet.append([None, "박새손", 70_000, "직장/기관", "", "참석", "계좌이체", 1, 0, "", "", "", "미발송"])
        workbook.save(exported)

        back = read_spreadsheet(exported).guests
        assert [g.name for g in back] == ["홍길동", "이영희", "김가족 & 김친지", "박새손"]
        assert [g.guest_id for g in back] == [1, 2, 3, 4]
        assert back[-1].payment is Payment.TRANSFER

    def test_cleared_review_column_means_reviewed(self, exported):
        _edit(exported, {(1, "확인필요"): "금액 후보가 여러 개입니다"})
        assert read_spreadsheet(exported).guests[0].needs_review
        _edit(exported, {(1, "확인필요"): None})
        assert not read_spreadsheet(exported).guests[0].needs_review

    def test_zero_amount_is_flagged(self, exported):
        _edit(exported, {(1, "축의금액"): None})
        assert WARN_NO_AMOUNT in read_spreadsheet(exported).guests[0].warnings

    def test_reads_list_even_if_saved_on_summary_sheet(self, exported):
        from openpyxl import load_workbook

        workbook = load_workbook(exported)
        workbook.active = workbook.sheetnames.index("정산 요약")
        workbook.save(exported)
        assert read_spreadsheet(exported).is_export

    def test_resaved_as_csv(self, exported, tmp_path):
        """엑셀에서 CSV로 다시 저장해 올려도 같은 형식으로 알아본다."""
        from openpyxl import load_workbook

        sheet = load_workbook(exported)[EXPORT_SHEET]
        path = tmp_path / "명단.csv"
        with path.open("w", encoding="utf-8-sig", newline="") as handle:
            csv.writer(handle).writerows(sheet.iter_rows(values_only=True))
        result = read_spreadsheet(path)
        assert result.is_export
        assert result.guests[0].sent_thanks is True


class TestDetection:
    def test_bank_statement_is_not_mistaken_for_export(self):
        rows = [["거래일시", "보낸분", "입금액", "메모"], ["2026-09-01", "홍길동", "50000", ""]]
        assert parse_export_rows(rows) is None
        assert parse_rows(rows)[0].name == "홍길동"

    def test_tolerates_reordered_and_dropped_columns(self):
        headers = [h for h in EXPORT_HEADERS if h not in ("순번", "감사메시지", "비고")]
        headers.reverse()
        row = {"성명": "홍길동", "축의금액": "100000", "관계분류": "친척/가족", "참석여부": "참석",
               "수령경로": "현금", "발송완료": "완료", "대인식권": "1", "소인식권": "0",
               "소속": "", "확인필요": ""}
        guests = parse_export_rows([headers, [row[h] for h in headers]])
        assert guests is not None
        assert guests[0].relation is Relation.FAMILY
        assert guests[0].sent_thanks is True


# ------------------------------------------------------------------ 화면


@pytest.fixture
def window(qt_app, monkeypatch):
    from PySide6.QtWidgets import QMessageBox

    from chugui.ui.main_window import MainWindow

    monkeypatch.setattr(QMessageBox, "question", lambda *a, **k: QMessageBox.StandardButton.Yes)
    win = MainWindow()
    win.show()
    qt_app.processEvents()
    yield win
    win.close()


class TestWindow:
    def test_reupload_replaces_instead_of_doubling(self, window, exported, guests):
        window._model.set_guests(parse_text(SOURCE))
        _edit(exported, {(1, "축의금액"): 150_000})
        window._load_file(str(exported))
        assert len(window._model.guests) == len(guests)
        assert window._model.guests[0].amount == 150_000
        assert not any(g.needs_review for g in window._model.guests)

    def test_reupload_can_be_undone(self, window, exported):
        window._model.set_guests(parse_text("박하객 1만원"))
        window._load_file(str(exported))
        window._undo_last()
        assert [g.name for g in window._model.guests] == ["박하객"]

    def test_declining_keeps_current_list(self, window, exported, monkeypatch):
        from PySide6.QtWidgets import QMessageBox

        window._model.set_guests(parse_text("박하객 1만원"))
        monkeypatch.setattr(QMessageBox, "question", lambda *a, **k: QMessageBox.StandardButton.No)
        window._load_file(str(exported))
        assert [g.name for g in window._model.guests] == ["박하객"]

    def test_other_spreadsheets_still_merge(self, window, tmp_path):
        path = tmp_path / "이체내역.csv"
        path.write_text("보낸분,입금액\n최이체,50000\n", encoding="utf-8-sig")
        window._model.set_guests(parse_text("박하객 1만원"))
        window._load_file(str(path))
        assert [g.name for g in window._model.guests] == ["박하객", "최이체"]
